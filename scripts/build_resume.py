"""Build the resume DOCX from the v2 markdown content.

Single-page A4 layout, Chinese-friendly font (微软雅黑 with Calibri fallback for ASCII).
Run from project root: python scripts/build_resume.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"D:/Projects/SyncBoard/刘佳个人简历_v2.docx")

CN_FONT = "微软雅黑"
EN_FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def set_run_font(run, size_pt: float, bold: bool = False, color: RGBColor | None = None):
    run.font.name = EN_FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)


def add_para(doc, text: str = "", *, size: float = 10.5, bold: bool = False,
             align=None, space_before: float = 0, space_after: float = 0,
             color: RGBColor | None = None, line_spacing: float = 1.15):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        set_run_font(run, size, bold=bold, color=color)
    return p


def add_h1(doc, text: str):
    p = add_para(doc, text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=0, space_after=2, color=ACCENT)
    return p


def add_section(doc, text: str):
    """Section header with bottom border."""
    p = add_para(doc, text, size=12, bold=True, space_before=8, space_after=2,
                 color=ACCENT)
    pPr = p._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_role_line(doc, left: str, right: str = ""):
    """Bold left, right pushed to right margin via tab."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.15
    from docx.shared import Cm
    pf.tab_stops.add_tab_stop(Cm(17.5), WD_ALIGN_PARAGRAPH.RIGHT)
    r1 = p.add_run(left)
    set_run_font(r1, 10.5, bold=True)
    if right:
        r2 = p.add_run("\t" + right)
        set_run_font(r2, 10.5, bold=False)
    return p


def add_bullet(doc, text: str, *, size: float = 10.0, bold_lead: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.2
    pf.left_indent = Cm(0.6)
    if bold_lead:
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size, bold=True)
        r2 = p.add_run(text)
        set_run_font(r2, size, bold=False)
    else:
        # parse **bold** markers inline
        parts = text.split("**")
        for i, chunk in enumerate(parts):
            if not chunk:
                continue
            run = p.add_run(chunk)
            set_run_font(run, size, bold=(i % 2 == 1))
    return p


def main() -> None:
    doc = Document()

    # Page setup — A4, tight margins for single-page fit
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    # Default style font
    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CN_FONT)

    # ===== Header =====
    add_h1(doc, "刘 佳")
    add_para(doc, "求职意向：测试工程师 / QA / 自动化测试",
             size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "电话：18623077084   ｜   邮箱：3420352607@qq.com   ｜   学号：6369685727710",
             size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    # ===== Education =====
    add_section(doc, "教育经历")
    add_role_line(doc, "重庆工商大学   |   自动化（本科）", "2022.9 – 2026.6")
    add_bullet(doc, "优秀学生综合奖学金二等奖、学习优秀奖学金、计算机趣味知识竞赛二等奖")

    # ===== Internship =====
    add_section(doc, "实习经历")
    add_role_line(doc, "西部创源智行科技（重庆）有限公司   |   实习测试工程师",
                  "2025.8 – 2026.2")
    intern_bullets = [
        "参与「乘用车智驾功能类人化测试方法研究与应用」课题项目与汽车网络安全项目，承担测试落地与缺陷收敛工作，推动版本按节点交付。",
        "**独立产出智驾「类人化」评测系统测试用例 500+**，覆盖系统业务逻辑与底层接口交互，提交评审一次通过率 ≥ 90%。",
        "使用 Postman + Linux（MobaXterm / SSH / vim）完成接口联调与日志排查，定位异步任务调度异常、算分逻辑偏差等多类问题，缩短开发回归周期。",
        "通过禅道闭环管理 bug 生命周期：创建 → 分配 → 复测 → 关闭，跟进修复进度并主导高优先级缺陷的复盘。",
        "输出系统操作手册、测试技术文档、企业级质量报告标准等过程资产，沉淀为团队可复用的测试规范。",
        "参与用例评审与回归策略制定，依据缺陷分布动态调整回归范围，提升迭代回归效率。",
    ]
    for b in intern_bullets:
        add_bullet(doc, b)

    # ===== Project =====
    add_section(doc, "项目经历")
    add_role_line(doc, "SyncBoard：项目协作 + 自动化测试一体化平台   |   全栈 + 测试负责人",
                  "2025.12 – 2026.2")
    add_para(doc,
             "基于 Django + Vue 3 自研的协作与质量平台，集成接口/UI 自动化、性能压测、CI/CD、AI 助手、RBAC 权限与实时同步；"
             "后端 5 个 Django 应用、44 个数据模型、约 90 个 REST 路由，前端 50 个组件 + 38 个路由的单页应用。",
             size=9.5, space_before=1, space_after=1, line_spacing=1.2)
    project_bullets = [
        "**接口自动化体系**：基于 pytest 沉淀 **299 条**自动化用例（25 个测试文件），覆盖鉴权、RBAC、看板、QA、缺陷工作流；自研统一断言引擎（JSONPath / 正则 / 表达式）+ 模板渲染（变量提取 → 链式传参），支持 Allure 报告与 CI 直出。",
        "**UI 自动化与录制回放**：基于 Playwright 实现浏览器录制器，支持 **18 类操作**（点击 / 双击 / 拖拽 / 文件上传 / iframe / 弹窗 / 滚动 / 截图等）+ **3 类断言**（文本 / 可见 / 存在），通过专用 WebSocket 把录制动作落成可编辑步骤库 +「一键试运行」回放，让非编码人员也能产出可维护的 E2E 用例。",
        "**性能压测**：基于 Locust 编写多场景脚本（登录 → 建项目 → 建任务 → 全文搜索 → 看板浏览），CI 中以 **50 并发用户 / ramp 5/s / 30s** 跑冒烟基线，建立性能回归卡点。",
        "**CI/CD 闭环**：用 GitHub Actions 搭建 **4 阶段并行流水线**（lint → API 测试 / E2E / 压测三路并行），叠加 Docker Buildx + Playwright 浏览器双层缓存；CD 实现滚动发布（scale=2 → 1）、dumpdata 自动备份、10 次重试健康检查、生产失败 **自动回滚 + Slack 告警**，保障每次合入可被回归脚本一键复现。",
        "**质量与安全治理**：主导一次全栈代码评审，输出 **30 项 P0–P3 缺陷清单**并闭环；修复 **8 个安全漏洞**（批量删除越权 IDOR、WebSocket XSS、提交前广播竞态、硬编码口令等），下沉 ProjectAccessMixin 统一权限校验与统一错误响应结构，新增 11 列数据库索引消除看板热路径 N+1，并补充 11 项头像上传安全测试。",
        "**实时同步链路**：基于 Django Channels + Redis 设计 4 个 WebSocket 通道（看板状态 / 聊天 / 全局通知 / QA 实时日志），前端封装单例 WebSocket（30s 心跳 + 3s 自动重连），通过 transaction.on_commit() 解决「先广播后落库」竞态，实时事件丢失率降至 0。",
        "**AI 提效**：接入 DeepSeek 并注册 **17 个 function-calling 工具**（创建 / 移动 / 查询任务、生成 API/UI 用例、触发流水线、拉取测试统计 / 失败用例等），把多步操作压缩为自然语言指令，演示场景下用例创建效率显著提升。",
    ]
    for b in project_bullets:
        add_bullet(doc, b, size=9.5)

    # ===== Skills =====
    add_section(doc, "相关技能")
    skills = [
        ("测试设计与执行：", "功能 / 接口 / UI / 性能 / 安全测试；等价类、边界值、场景法、状态迁移；缺陷生命周期管理（禅道）"),
        ("自动化框架：", "pytest、Playwright、Postman、Locust、Allure"),
        ("编程语言：", "Python（主）、SQL、JavaScript / TypeScript、C"),
        ("全栈开发：", "Django / DRF / Channels、Vue 3 / Pinia、WebSocket、Celery"),
        ("数据库 / 中间件：", "MySQL、Redis、Elasticsearch"),
        ("DevOps / 运维：", "GitHub Actions、Docker / Docker Compose、Linux / SSH / MobaXterm、日志排查"),
    ]
    for lead, rest in skills:
        add_bullet(doc, rest, size=10, bold_lead=lead)

    # ===== Self assessment =====
    add_section(doc, "自我评价")
    add_para(doc,
             "应届生在校期间通过实习与自研项目，积累了从用例设计、接口/UI 自动化到 CI/CD、日志排查的一线实践，"
             "养成了写完即测、出问题先看日志再改代码的习惯。学习能力较强，遇到不熟悉的栈愿意翻文档与源码逐层定位，"
             "不会把问题甩回上游；性格稳、沟通顺畅，能配合开发与产品对齐细节，也乐于把跑通的流程沉淀成文档供同事复用。"
             "期待加入团队后，从测试基础工作做起，逐步承担更系统的质量保障职责。",
             size=10, space_before=2, space_after=0, line_spacing=1.3)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
